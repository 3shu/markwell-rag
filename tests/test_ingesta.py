"""Pruebas unitarias de la extracción y división por artículo (sin AWS ni base de datos)."""

import io
import zipfile

from docx import Document

from markwell_rag.ingesta import (
    MAX_CARACTERES_RECOMENDADO,
    Articulo,
    dividir_por_articulo,
    extraer_parrafos_docx,
    parrafos_del_cuerpo,
    validar,
)


def etiquetas(division):
    return [a.etiqueta for a in division.articulos]


def test_agrupa_parrafos_por_articulo():
    division = dividir_por_articulo(
        [
            "Artículo 1. Objeto. Este reglamento regula la atención de reclamos.",
            "Artículo 2. Plazo. La entidad debe responder en diez días.",
            "El plazo se cuenta desde la recepción del reclamo.",
        ]
    )

    assert etiquetas(division) == ["Art. 1", "Art. 2"]
    assert division.articulos[1].texto == (
        "Artículo 2. Plazo. La entidad debe responder en diez días.\n"
        "El plazo se cuenta desde la recepción del reclamo."
    )


def test_separa_el_preambulo():
    division = dividir_por_articulo(
        [
            "El Consejo Nacional de Supervisión del Sistema Financiero,",
            "considerando que:",
            "Artículo 1. Objeto.",
        ]
    )

    assert division.preambulo == [
        "El Consejo Nacional de Supervisión del Sistema Financiero,",
        "considerando que:",
    ]
    assert etiquetas(division) == ["Art. 1"]


def test_reconoce_variantes_de_encabezado():
    division = dividir_por_articulo(
        [
            "ARTÍCULO 1.- Objeto.",
            "Articulo 2: Alcance.",
            "Artículo 3º.- Definiciones.",
            "Artículo 3 bis. Excepciones.",
        ]
    )

    assert etiquetas(division) == ["Art. 1", "Art. 2", "Art. 3", "Art. 3 bis"]


def test_una_referencia_interna_no_abre_un_articulo():
    division = dividir_por_articulo(
        [
            "Artículo 1. Objeto.",
            "De conformidad con el artículo 25 de la Ley 8653, la entidad informará.",
        ]
    )

    assert etiquetas(division) == ["Art. 1"]


def test_descarta_encabezados_de_capitulo():
    division = dividir_por_articulo(
        [
            "Artículo 1. Objeto.",
            "TÍTULO II",
            "ATENCIÓN DE AVISOS DE SINIESTRO",
            "CAPÍTULO ÚNICO",
            "Del procedimiento de aviso de siniestro",
            "Artículo 2.Vigencia.",
        ]
    )

    assert division.articulos[0].texto == "Artículo 1. Objeto."
    assert etiquetas(division) == ["Art. 1", "Art. 2"]


def test_reconoce_transitorios():
    division = dividir_por_articulo(
        [
            "Artículo 1. Vigencia.",
            "Transitorio I. Las entidades tendrán seis meses para ajustarse.",
            "TRANSITORIO único.- Plazo de adecuación.",
        ]
    )

    assert etiquetas(division) == ["Art. 1", "Transitorio I", "Transitorio único"]


def test_validar_informa_saltos_repetidos_y_largos():
    articulos = [
        Articulo("Art. 1", "uno"),
        Articulo("Art. 3", "tres"),
        Articulo("Art. 3", "x" * (MAX_CARACTERES_RECOMENDADO + 1)),
    ]

    avisos = validar(articulos)

    assert "Salto en la numeración: Art. 1 seguido de Art. 3" in avisos
    assert "Etiqueta repetida: Art. 3" in avisos
    assert any("caracteres" in a for a in avisos)


def test_extrae_parrafos_de_un_docx(tmp_path):
    ruta = tmp_path / "reglamento.docx"
    documento = Document()
    documento.add_paragraph("Artículo 1. Objeto.")
    documento.add_paragraph("   ")  # párrafo vacío: se descarta
    documento.add_paragraph("  Artículo 2. Alcance.  ")
    documento.save(str(ruta))

    assert extraer_parrafos_docx(ruta) == [
        "Artículo 1. Objeto.",
        "Artículo 2. Alcance.",
    ]


def test_elige_el_documento_incrustado_con_articulos(tmp_path):
    def docx_en_bytes(*parrafos: str) -> bytes:
        documento = Document()
        for p in parrafos:
            documento.add_paragraph(p)
        datos = io.BytesIO()
        documento.save(datos)
        return datos.getvalue()

    # Contenedor como el de SUGESE: el articulado viene incrustado, no en la página principal.
    ruta = tmp_path / "contenedor.docx"
    ruta.write_bytes(docx_en_bytes("ACUERDO SUGESE 99-99", "Reglamento"))
    with zipfile.ZipFile(ruta, "a") as paquete:
        paquete.writestr(
            "word/embeddings/Microsoft_Word_Document.docx",
            docx_en_bytes("Considerando que:"),
        )
        paquete.writestr(
            "word/embeddings/Microsoft_Word_Document1.docx",
            docx_en_bytes("Artículo 1.Del objeto", "Artículo 2.Ámbito"),
        )

    nombre, parrafos = parrafos_del_cuerpo(ruta)

    assert nombre == "word/embeddings/Microsoft_Word_Document1.docx"
    assert parrafos == ["Artículo 1.Del objeto", "Artículo 2.Ámbito"]
